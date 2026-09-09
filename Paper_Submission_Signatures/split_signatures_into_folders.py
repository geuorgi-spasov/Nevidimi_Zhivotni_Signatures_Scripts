"""
Splits a raw CSV export of signatures into multiple Word documents for
paper submission. Each document contains ROWS_PER_FILE signatures
(default: 1000), laid out ROWS_PER_PAGE per landscape page (default: 10),
with a footer that page-numbers continuously across all files.

All generated .docx files are written into OUTPUT_DOCX_FOLDER, which is
created automatically if it doesn't exist.

Usage:
    python split_signatures_into_folders.py

See README.md for the full workflow.
"""

from __future__ import annotations

import os
import time

import pandas as pd
from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT
from docx.enum.table import (
    WD_ALIGN_VERTICAL,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table

# ---------------------------------------------------------------------------
# Configuration — change these if your file names or layout differ
# ---------------------------------------------------------------------------

INPUT_CSV = "Signatures_from_the_database_raw.csv"
OUTPUT_DOCX_FOLDER = "signatures_docx"

# Grouping
ROWS_PER_FILE = 1000   # signatures per submission folder/document
ROWS_PER_PAGE = 10     # data rows per printed page

# Page layout (landscape US Letter)
PAGE_WIDTH_IN = 11.0
PAGE_HEIGHT_IN = 8.5
LEFT_MARGIN_IN = 0.75
RIGHT_MARGIN_IN = 0.75
TOP_MARGIN_IN = 0.75
BOTTOM_MARGIN_IN = 1.0

# Row layout
ROW_HEIGHT_IN = 0.52
BODY_FONT = "Arial"
FOOTER_FONT = "Cambria"
BODY_FONT_SIZE_PT = 12

# Footer text (Bulgarian)
ORGANIZATION_NAME = 'Сдружение „Невидими животни"'

# Column widths in inches, tuned for a 6-column CSV. The numbers are scaled
# so their total exactly fills the available page width.
DEFAULT_COLUMN_WIDTHS_IN = [0.70, 1.36, 1.45, 3.44, 1.31, 1.19]

# Encodings and separators tried when auto-detecting the CSV format
CSV_SEPARATORS = [",", ";", "\t", "|"]
CSV_ENCODINGS = ["utf-8", "latin-1", "cp1252", "iso-8859-1"]


# ---------------------------------------------------------------------------
# Utilities
# ---------------------------------------------------------------------------

def _format_duration(seconds: float) -> str:
    """Format seconds as ``'Xs'`` or ``'Xm Ys'``."""
    if seconds < 60:
        return f"{seconds:.0f}s"
    return f"{int(seconds) // 60}m {int(seconds) % 60}s"


# ---------------------------------------------------------------------------
# CSV reading
# ---------------------------------------------------------------------------

def _try_read_csv(path: str, sep: str, encoding: str = "utf-8") -> pd.DataFrame | None:
    """Return the DataFrame only if it parses into more than one column."""
    try:
        df = pd.read_csv(path, sep=sep, encoding=encoding)
    except Exception:
        return None
    if len(df.columns) > 1:
        return df
    return None


def read_signatures_csv(path: str) -> pd.DataFrame:
    """Read the raw CSV, trying common separators and encodings."""
    # 1) Try standard separators with utf-8
    for sep in CSV_SEPARATORS:
        df = _try_read_csv(path, sep)
        if df is not None:
            print(f"Read '{path}' using separator {sep!r} and utf-8.")
            return df

    # 2) Pick the most common separator on the first line
    with open(path, "r", encoding="utf-8") as fh:
        first_line = fh.readline().strip()
    best_sep = max(CSV_SEPARATORS, key=first_line.count)
    if first_line.count(best_sep) > 0:
        df = _try_read_csv(path, best_sep)
        if df is not None:
            print(f"Read '{path}' using detected separator {best_sep!r}.")
            return df

    # 3) Fall back to alternative encodings
    for encoding in CSV_ENCODINGS:
        for sep in CSV_SEPARATORS:
            df = _try_read_csv(path, sep, encoding)
            if df is not None:
                print(f"Read '{path}' using {sep!r} and encoding '{encoding}'.")
                return df

    raise ValueError(f"Could not parse '{path}' as a multi-column CSV.")


# ---------------------------------------------------------------------------
# Document building
# ---------------------------------------------------------------------------

def scale_column_widths(num_columns: int) -> list[float]:
    """Return column widths in inches that fit the available page width."""
    available_in = PAGE_WIDTH_IN - LEFT_MARGIN_IN - RIGHT_MARGIN_IN

    if num_columns == len(DEFAULT_COLUMN_WIDTHS_IN):
        widths = DEFAULT_COLUMN_WIDTHS_IN
    elif num_columns < len(DEFAULT_COLUMN_WIDTHS_IN):
        widths = DEFAULT_COLUMN_WIDTHS_IN[:num_columns]
    else:
        # More columns than we have presets for: distribute evenly
        return [available_in / num_columns] * num_columns

    scale = available_in / sum(widths)
    return [w * scale for w in widths]


def apply_full_table_borders(table: Table) -> None:
    """Draw visible single-line borders on every cell of the table."""
    border_style = {
        qn("w:val"): "single",
        qn("w:sz"): "4",
        qn("w:space"): "0",
        qn("w:color"): "000000",
    }

    # Table-level borders (outer + inner)
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        for attr, value in border_style.items():
            edge.set(attr, value)
        tbl_borders.append(edge)
    tbl_pr.append(tbl_borders)

    # Cell-level borders (belt and braces — some viewers need both)
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_borders = OxmlElement("w:tcBorders")
            for name in ("top", "left", "bottom", "right"):
                edge = OxmlElement(f"w:{name}")
                for attr, value in border_style.items():
                    edge.set(attr, value)
                tc_borders.append(edge)
            tc_pr.append(tc_borders)


def create_landscape_document() -> DocumentType:
    """Create a new Word document set up for landscape printing."""
    doc = Document()
    doc.styles["Normal"].font.name = BODY_FONT

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.left_margin = Inches(LEFT_MARGIN_IN)
    section.right_margin = Inches(RIGHT_MARGIN_IN)
    section.top_margin = Inches(TOP_MARGIN_IN)
    section.bottom_margin = Inches(BOTTOM_MARGIN_IN)

    return doc


def add_page_table(
    doc: DocumentType,
    header: list[str],
    page_rows: pd.DataFrame,
    column_widths: list[float],
) -> None:
    """Add one page's worth of data as a bordered table."""
    table = doc.add_table(rows=len(page_rows) + 1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i, width in enumerate(column_widths):
        if i < len(table.columns):
            table.columns[i].width = Inches(width)

    # Header row: bold, centered, body font size
    for i, col_name in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = str(col_name)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(BODY_FONT_SIZE_PT)

    # Data rows: left-aligned, body font size
    for row_idx, (_, data_row) in enumerate(page_rows.iterrows()):
        table_row = table.rows[row_idx + 1]
        for col_idx, value in enumerate(data_row):
            if col_idx >= len(table_row.cells):
                continue
            cell = table_row.cells[col_idx]
            cell.text = str(value) if pd.notna(value) else ""
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(BODY_FONT_SIZE_PT)

    # Fixed row height with vertically centered cells
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        row.height = Inches(ROW_HEIGHT_IN)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    apply_full_table_borders(table)


def add_page_footer(
    doc: DocumentType, global_page_number: int, folder_number: int
) -> None:
    """Write the per-page footer: blank line, page/folder info, org name."""
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.add_run().add_break()  # blank line before the footer

    info_run = footer.add_run(
        f"Стр. {global_page_number}, папка {folder_number}"
    )
    info_run.font.name = FOOTER_FONT
    info_run.font.size = Pt(BODY_FONT_SIZE_PT)
    info_run.italic = True
    info_run.add_break()

    org_run = footer.add_run(ORGANIZATION_NAME)
    org_run.font.name = FOOTER_FONT
    org_run.font.size = Pt(BODY_FONT_SIZE_PT)
    org_run.italic = True


def build_folder_document(
    folder_number: int,
    chunk: pd.DataFrame,
    header: list[str],
    column_widths: list[float],
) -> tuple[DocumentType, int]:
    """Build the Word document for a single submission folder (~1000 rows)."""
    doc = create_landscape_document()

    pages_in_file = (len(chunk) + ROWS_PER_PAGE - 1) // ROWS_PER_PAGE
    # Pages in a *full* folder — used so page numbering continues across files
    pages_per_full_folder = (ROWS_PER_FILE + ROWS_PER_PAGE - 1) // ROWS_PER_PAGE

    for page_num in range(1, pages_in_file + 1):
        start = (page_num - 1) * ROWS_PER_PAGE
        end = min(start + ROWS_PER_PAGE, len(chunk))
        page_rows = chunk.iloc[start:end]

        add_page_table(doc, header, page_rows, column_widths)

        global_page = page_num + (folder_number - 1) * pages_per_full_folder
        add_page_footer(doc, global_page, folder_number)

        if page_num < pages_in_file:
            doc.add_page_break()

    return doc, pages_in_file


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def main() -> None:
    print(f"Reading '{INPUT_CSV}'...")
    try:
        df = read_signatures_csv(INPUT_CSV)
    except FileNotFoundError:
        print(f"ERROR: file '{INPUT_CSV}' not found in this folder.")
        return
    except ValueError as exc:
        print(f"ERROR: {exc}")
        return

    print(f"Read {len(df)} rows with {len(df.columns)} columns.")
    print(f"Columns: {list(df.columns)}")

    header = df.columns.tolist()
    column_widths = scale_column_widths(len(header))
    print(f"Column widths (inches): {[round(w, 2) for w in column_widths]}")

    os.makedirs(OUTPUT_DOCX_FOLDER, exist_ok=True)
    print(f"Writing output into '{OUTPUT_DOCX_FOLDER}/'.")

    total_files = (len(df) + ROWS_PER_FILE - 1) // ROWS_PER_FILE
    print(f"\nCreating {total_files} submission file(s)...\n")

    overall_start = time.perf_counter()

    for folder_number in range(1, total_files + 1):
        start = (folder_number - 1) * ROWS_PER_FILE
        end = min(start + ROWS_PER_FILE, len(df))
        chunk = df.iloc[start:end]

        first_id = str(chunk.iloc[0, 0])
        last_id = str(chunk.iloc[-1, 0])

        doc, pages = build_folder_document(
            folder_number, chunk, header, column_widths
        )
        filename = (
            f"Папка {folder_number} с подписи от {first_id} до {last_id}.docx"
        )
        output_path = os.path.join(OUTPUT_DOCX_FOLDER, filename)
        doc.save(output_path)

        elapsed = time.perf_counter() - overall_start
        avg_per_folder = elapsed / folder_number
        remaining = avg_per_folder * (total_files - folder_number)
        print(
            f"  [{folder_number}/{total_files}] '{filename}' — "
            f"{len(chunk)} signatures, {pages} pages"
            + (
                f"  (~{_format_duration(remaining)} remaining)"
                if folder_number < total_files
                else ""
            )
        )

    total_time = time.perf_counter() - overall_start
    print(
        f"\nDone in {_format_duration(total_time)}. "
        f"Created {total_files} file(s) in '{OUTPUT_DOCX_FOLDER}/'."
    )


if __name__ == "__main__":
    main()
